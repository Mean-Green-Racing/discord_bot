"""Read-only local Git mirror and inexpensive AGENTS.md-guided document search.

GitHub and local search do not use Anthropic tokens. Only selected excerpts
are passed to the caller to include in a Claude request.
"""
from __future__ import annotations

import asyncio
from collections import Counter, defaultdict
from dataclasses import dataclass
import hashlib
import json
import math
import os
from pathlib import Path, PurePosixPath
import re
import shlex
import stat
import subprocess
import tempfile
import threading
import time
from urllib.parse import quote, unquote


FILE_TYPES = {".md", ".markdown", ".txt", ".rst", ".csv", ".yaml", ".yml", ".json", ".pdf", ".docx"}
MAX_FILE_BYTES = 3_000_000
MAX_DOC_CHARS = 7500
MAX_PASSAGES = 5
MAX_FILES = 3
# Only broaden retrieval on questions that actually need cross-document reasoning.
DEEP_MAX_PASSAGES = 8
DEEP_MAX_FILES = 5
DEEP_MAX_DOC_CHARS = 11500
STOPWORDS = {
    "a", "an", "and", "are", "as", "at", "be", "by", "can", "could", "do", "does",
    "for", "from", "how", "i", "in", "is", "it", "me", "my", "of", "on", "or",
    "our", "please", "should", "tell", "that", "the", "their", "this", "to", "we",
    "what", "when", "where", "which", "who", "why", "will", "with", "would", "you",
}
PATH_RE = re.compile(
    r"(?<![\w./-])(?:\.\.?/)?[\w./-]+\.(?:md|markdown|txt|rst|csv|yaml|yml|json|pdf|docx)\b",
    re.IGNORECASE,
)
HEADING_RE = re.compile(r"^\s{0,3}#{1,6}\s+(.+)$")
WORD_RE = re.compile(r"[a-z0-9]+", re.IGNORECASE)


def words(value: str) -> list[str]:
    result = []
    for term in WORD_RE.findall(value.lower()):
        if term in STOPWORDS:
            continue
        # A small, deterministic normalization step; no external embedding API.
        if term.endswith("ies") and len(term) > 5:
            term = term[:-3] + "y"
        elif term.endswith("s") and not term.endswith("ss") and len(term) > 4:
            term = term[:-1]
        result.append(term)
    return result


def github_repo_slug(url: str) -> str | None:
    """Return owner/repo for an HTTPS or SSH GitHub remote, if recognized."""
    url = url.strip()
    for prefix in ("https://github.com/", "git@github.com:", "ssh://git@github.com/"):
        if url.startswith(prefix):
            slug = url[len(prefix):].rstrip("/")
            if slug.endswith(".git"):
                slug = slug[:-4]
            if re.fullmatch(r"[A-Za-z0-9_.-]+/[A-Za-z0-9_.-]+", slug):
                return slug
            raise ValueError(f"Invalid GitHub repository URL: {url}")
    return None


def ssh_repo_url(url: str) -> str:
    """For GitHub, force SSH even if an old .env still specifies HTTPS."""
    slug = github_repo_slug(url)
    if slug is not None:
        return f"git@github.com:{slug}.git"
    return url


@dataclass(frozen=True)
class Passage:
    path: str
    start: int
    end: int
    content: str
    title: str = ""
    source_kind: str = "lines"  # lines, page, paragraphs

    @property
    def location(self) -> str:
        if self.source_kind == "page":
            return f"page {self.start}"
        if self.source_kind == "paragraphs":
            return f"paragraphs {self.start}-{self.end}"
        return f"lines {self.start}-{self.end}"


@dataclass(frozen=True)
class SearchHit:
    passage: Passage
    score: float


class KnowledgeBase:
    def __init__(
        self,
        folder: Path,
        repo_url: str,
        branch: str = "main",
        ssh_key_path: str | Path = "",
        refresh_seconds: int = 600,
        semantic_enabled: bool = True,
        embedding_model_name: str = "BAAI/bge-small-en-v1.5",
        embedding_threads: int = 1,
        embedding_batch_size: int = 8,
    ) -> None:
        self.folder = folder
        self.repo_url = ssh_repo_url(repo_url)
        self.branch = branch
        self.ssh_key_path = str(ssh_key_path).strip()
        self.refresh_seconds = max(60, refresh_seconds)
        self._lock = asyncio.Lock()
        self._checked_at = 0.0
        self.commit = ""
        self.last_error = ""
        self.passages: list[Passage] = []
        self.routes: dict[str, str] = {}
        self.doc_words: dict[str, set[str]] = {}
        self.file_count = 0
        # Local FastEmbed/ONNX inference, not a paid API. Keep the cache outside
        # the disposable Git mirror, which is reset on every repository update.
        self.semantic_enabled = semantic_enabled
        self.embedding_model_name = embedding_model_name
        self.embedding_threads = max(1, int(embedding_threads))
        self.embedding_batch_size = min(16, max(1, int(embedding_batch_size)))
        self.embedding_cache_dir = self.folder.parent / "embedding_models"
        self.embedding_index_path = self.folder.parent / "semantic_index.npz"
        self.semantic_status = "disabled" if not semantic_enabled else "pending"
        self.semantic_last_error = ""
        self._semantic_signature = ""
        self._failed_signature = ""
        self._vectors_signature = ""
        self._semantic_vectors = None
        self._semantic_task: asyncio.Task | None = None
        self._embedding_model = None
        self._encoder_lock = threading.RLock()

    def _private_key(self) -> Path:
        """Validate the *local private key* before any network access."""
        if not self.ssh_key_path:
            raise RuntimeError(
                "Set GITHUB_DEPLOY_KEY_PATH in .env to the LOCAL PRIVATE key file "
                "matching the repository's GitHub deploy key (not the .pub file)."
            )
        key = Path(os.path.expandvars(self.ssh_key_path)).expanduser()
        if not key.is_absolute():
            key = Path(__file__).resolve().parent / key
        key = key.resolve()
        if key.suffix == ".pub":
            raise RuntimeError(f"Deploy key path must point to the PRIVATE key, not: {key}")
        if not key.is_file():
            raise RuntimeError(f"Deploy private key not found: {key}")
        if os.name == "posix" and (stat.S_IMODE(key.stat().st_mode) & 0o077):
            raise RuntimeError(
                f"Deploy private key has unsafe permissions: {key}. Run: chmod 600 '{key}'"
            )
        return key

    def _git(self, *args: str, cwd: Path | None = None) -> str:
        env = os.environ.copy()
        env["GIT_TERMINAL_PROMPT"] = "0"
        # Never include the private key contents in .env or in a Git URL.
        # Only clone/fetch access the network. Other commands work offline.
        if args and args[0] in {"clone", "fetch", "ls-remote"} and github_repo_slug(self.repo_url):
            key = self._private_key()
            env["GIT_SSH_COMMAND"] = shlex.join([
                "ssh", "-i", str(key),
                "-o", "IdentitiesOnly=yes",
                "-o", "BatchMode=yes",
                "-o", "StrictHostKeyChecking=yes",
                "-o", "ConnectTimeout=10",
            ])
        try:
            result = subprocess.run(
                ["git", *args], cwd=cwd, env=env,
                capture_output=True, text=True, timeout=75, check=False,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            raise RuntimeError(f"Git unavailable or timed out: {exc}") from exc
        if result.returncode:
            error = (result.stderr or result.stdout or "unknown error").strip()
            if "Host key verification failed" in error:
                error += " | Trust github.com after verifying its official SSH host-key fingerprint."
            elif "Permission denied (publickey)" in error:
                error += " | Verify that the matching PUBLIC key is a read-only deploy key on this repo."
            raise RuntimeError(f"git {' '.join(args[:2])} failed: {error[-550:]}")
        return result.stdout.strip()

    def _ensure_correct_origin(self) -> None:
        """Safely migrate an existing checkout of the same repo from HTTPS to SSH."""
        origin = self._git("remote", "get-url", "origin", cwd=self.folder)
        if origin == self.repo_url:
            return
        old_slug = github_repo_slug(origin)
        new_slug = github_repo_slug(self.repo_url)
        if old_slug and new_slug and old_slug.lower() == new_slug.lower():
            self._git("remote", "set-url", "origin", self.repo_url, cwd=self.folder)
            return
        raise RuntimeError(
            "Existing local clone has an unexpected Git remote (not the configured KB). "
            "Refusing to change it automatically; inspect the origin before continuing."
        )

    def _sync(self) -> None:
        self.folder.parent.mkdir(parents=True, exist_ok=True)
        if not (self.folder / ".git").is_dir():
            if self.folder.exists() and any(self.folder.iterdir()):
                raise RuntimeError(f"Cannot clone into non-empty folder: {self.folder}")
            self._git(
                "clone", "--quiet", "--depth", "1", "--single-branch", "--branch",
                self.branch, self.repo_url, str(self.folder),
            )
        else:
            # This is a dedicated, disposable *mirror*; do not edit it locally.
            self._ensure_correct_origin()
            self._git("fetch", "--quiet", "--depth", "1", "origin", self.branch, cwd=self.folder)
            self._git("reset", "--quiet", "--hard", "FETCH_HEAD", cwd=self.folder)
        current = self._git("rev-parse", "HEAD", cwd=self.folder)
        if current != self.commit or not self.passages:
            self._index_folder()
            self.commit = current

    async def ensure_ready(self, force: bool = False) -> None:
        if (
            not force and self.passages
            and time.monotonic() - self._checked_at < self.refresh_seconds
        ):
            self._schedule_semantic_build()
            return
        async with self._lock:
            if (
                not force and self.passages
                and time.monotonic() - self._checked_at < self.refresh_seconds
            ):
                self._schedule_semantic_build()
                return
            self._checked_at = time.monotonic()
            try:
                await asyncio.to_thread(self._sync)
                self.last_error = ""
            except Exception as exc:
                self.last_error = str(exc)
                # If the network is down, reuse the last successfully indexed
                # local clone, but clearly report that it may be stale.
                if not self.passages and self.folder.exists():
                    try:
                        self._index_folder()
                        self.commit = self._git("rev-parse", "HEAD", cwd=self.folder)
                    except Exception:
                        pass
                if not self.passages:
                    raise RuntimeError(self.last_error) from exc
            # Never hold up Discord answers while the ONNX model downloads or
            # the index is built. Plain keyword search is immediately usable.
            self._schedule_semantic_build()

    def _schedule_semantic_build(self) -> None:
        if not self.semantic_enabled or not self.passages:
            return
        if self._vectors_signature == self._semantic_signature and self._semantic_vectors is not None:
            return
        if self._semantic_task is not None and not self._semantic_task.done():
            return
        # Don't retry failed model downloads on every chat message; the next
        # bot restart can retry, or a successful repository update may retry.
        if self.semantic_status == "unavailable" and self._failed_signature == self._semantic_signature:
            return
        self.semantic_status = "building"
        self._semantic_task = asyncio.create_task(self._build_semantic_async())

    async def _build_semantic_async(self) -> None:
        # Snapshot immutable Passage objects so a Git refresh cannot cause an
        # index to be attached to a different version of the repository.
        signature = self._semantic_signature
        passages = tuple(self.passages)
        try:
            await asyncio.to_thread(self._build_semantic_sync, signature, passages)
        except Exception as exc:
            if self._semantic_signature == signature:
                self._failed_signature = signature
                self.semantic_status = "unavailable"
                self.semantic_last_error = f"{type(exc).__name__}: {str(exc)[:260]}"
                print(f"KB semantic search unavailable (keyword fallback active): {self.semantic_last_error}")
        else:
            if self._semantic_signature == signature:
                self.semantic_status = "ready"
                self.semantic_last_error = ""
        finally:
            if self._semantic_signature != signature:
                # Git updated while the older document set was embedding.
                self._semantic_task = None
                self._schedule_semantic_build()

    def _get_embedding_model(self):
        with self._encoder_lock:
            if self._embedding_model is None:
                # Lazy import keeps keyword-only mode functional even if the
                # optional FastEmbed package isn't installed / CPU is unsupported.
                from fastembed import TextEmbedding
                self.embedding_cache_dir.mkdir(parents=True, exist_ok=True)
                self._embedding_model = TextEmbedding(
                    model_name=self.embedding_model_name,
                    cache_dir=str(self.embedding_cache_dir),
                    threads=self.embedding_threads,
                    providers=["CPUExecutionProvider"],
                )
            return self._embedding_model

    @staticmethod
    def _normalize_vectors(vectors):
        import numpy as np
        array = np.asarray(vectors, dtype=np.float32)
        if array.ndim != 2 or not len(array) or not np.isfinite(array).all():
            raise ValueError("Invalid embeddings returned from model")
        lengths = np.linalg.norm(array, axis=1, keepdims=True)
        return np.ascontiguousarray(array / np.maximum(lengths, 1e-12))

    def _build_semantic_sync(self, signature: str, passages: tuple[Passage, ...]) -> None:
        # Single ONNX encoder for document indexing and interactive queries.
        # 1 thread + 8-document batches lowers memory peaks on a 1 GB VPS.
        with self._encoder_lock:
            model = self._get_embedding_model()
            import numpy as np
            inputs = [f"{p.path}\n{p.title}\n{p.content}" for p in passages]
            embeddings = list(model.embed(inputs, batch_size=self.embedding_batch_size))
            if len(embeddings) != len(passages):
                raise ValueError("Embedding count didn't match document count")
            vectors = self._normalize_vectors(embeddings)
            if self._semantic_signature != signature:
                return
            self._semantic_vectors = vectors
            self._vectors_signature = signature
            # Write the reusable local index atomically, without pickle.
            self.embedding_index_path.parent.mkdir(parents=True, exist_ok=True)
            temp_path = None
            try:
                with tempfile.NamedTemporaryFile(
                    prefix="semantic-", suffix=".npz",
                    dir=self.embedding_index_path.parent, delete=False
                ) as handle:
                    temp_path = Path(handle.name)
                    np.savez_compressed(handle, vectors=vectors, signature=np.array(signature),
                                        model=np.array(self.embedding_model_name))
                if self._semantic_signature == signature:
                    temp_path.replace(self.embedding_index_path)
            finally:
                if temp_path is not None:
                    temp_path.unlink(missing_ok=True)

    def _load_semantic_cache(self) -> bool:
        if not self.semantic_enabled or not self.embedding_index_path.exists():
            return False
        import numpy as np
        try:
            with np.load(self.embedding_index_path, allow_pickle=False) as cache:
                if (cache["signature"].item() != self._semantic_signature
                        or cache["model"].item() != self.embedding_model_name):
                    return False
                vectors = cache["vectors"]
                if (vectors.ndim != 2 or vectors.shape[0] != len(self.passages)
                        or vectors.shape[1] > 2048 or not np.isfinite(vectors).all()):
                    return False
                self._semantic_vectors = np.ascontiguousarray(vectors, dtype=np.float32)
                self._vectors_signature = self._semantic_signature
                self.semantic_status = "ready"
                return True
        except (KeyError, ValueError, OSError) as exc:
            print(f"Invalid cached semantic index; rebuilding: {type(exc).__name__}")
        return False

    def _index_folder(self) -> None:
        all_files: dict[str, Path] = {}
        for path in self.folder.rglob("*"):
            if not path.is_file() or path.is_symlink():
                continue
            rel = path.relative_to(self.folder)
            if any(part.startswith(".") for part in rel.parts):
                continue
            if path.suffix.lower() not in FILE_TYPES:
                continue
            if path.stat().st_size > MAX_FILE_BYTES:
                continue
            all_files[rel.as_posix()] = path

        route_lines: dict[str, list[str]] = defaultdict(list)
        for rel, absolute in all_files.items():
            if absolute.name.lower() != "agents.md":
                continue
            guide = absolute.read_text(encoding="utf-8", errors="replace")
            heading = ""
            for line in guide.splitlines():
                heading_match = HEADING_RE.match(line)
                if heading_match:
                    heading = heading_match.group(1)
                for raw_ref in PATH_RE.findall(unquote(line)):
                    # Guides may mention ./foo.md or docs/foo.md; resolve
                    # relative to AGENTS.md first, then repository root.
                    for candidate in (
                        (PurePosixPath(rel).parent / raw_ref).as_posix(),
                        PurePosixPath(raw_ref).as_posix(),
                    ):
                        normalized = os.path.normpath(candidate).replace("\\", "/")
                        if normalized in all_files:
                            route_lines[normalized].append(f"{heading} {line.strip()[:240]}")
                            break

        passages: list[Passage] = []
        for rel, path in sorted(all_files.items()):
            try:
                if path.suffix.lower() == ".pdf":
                    from pypdf import PdfReader
                    reader = PdfReader(str(path))
                    for page_number, page in enumerate(reader.pages[:80], 1):
                        extracted = page.extract_text() or ""
                        if not extracted.strip():
                            continue
                        passages.extend(
                            self._make_passages(rel, extracted.splitlines(), "page", page_number)
                        )
                elif path.suffix.lower() == ".docx":
                    from docx import Document
                    document = Document(str(path))
                    lines = [p.text for p in document.paragraphs]
                    for table in document.tables:
                        for row in table.rows:
                            lines.append(" | ".join(cell.text for cell in row.cells))
                    passages.extend(self._make_passages(rel, lines, "paragraphs"))
                else:
                    raw = path.read_text(encoding="utf-8", errors="replace")
                    passages.extend(self._make_passages(rel, raw.splitlines(), "lines"))
            except Exception as exc:
                print(f"Index warning: skipped {rel}: {type(exc).__name__}: {exc}")

        if not passages:
            raise RuntimeError("No readable documents in cloned repo (checked Markdown, text, PDF, DOCX).")
        self.passages = passages
        self.routes = {key: " ".join(value)[:2500] for key, value in route_lines.items()}
        self.file_count = len({p.path for p in passages})
        doc_words: dict[str, set[str]] = defaultdict(set)
        for passage in passages:
            doc_words[passage.path].update(words(passage.content))
            doc_words[passage.path].update(words(passage.title))
        for key in doc_words:
            doc_words[key].update(words(key))
            doc_words[key].update(words(self.routes.get(key, "")))
        self.doc_words = dict(doc_words)
        digest = hashlib.sha256()
        digest.update(b"mgr-semantic-v1\0")
        digest.update(self.embedding_model_name.encode("utf-8"))
        for p in passages:
            digest.update(json.dumps(
                [p.path, p.start, p.end, p.title, p.source_kind, p.content],
                ensure_ascii=False, separators=(",", ":")
            ).encode("utf-8"))
        signature = digest.hexdigest()
        if signature != self._semantic_signature:
            self._semantic_signature = signature
            self._semantic_vectors = None
            self._vectors_signature = ""
            self.semantic_status = "pending" if self.semantic_enabled else "disabled"
            self._load_semantic_cache()

    @staticmethod
    def _make_passages(path: str, lines: list[str], kind: str, page: int = 0) -> list[Passage]:
        chunks = []
        buffer: list[tuple[int, str]] = []
        n = 0
        title = ""
        for number, line in enumerate(lines, 1):
            heading_match = HEADING_RE.match(line)
            if heading_match:
                title = heading_match.group(1)
            # Split individual long lines, not entire files into one massive chunk.
            fragments = [line[i:i + 700] for i in range(0, len(line), 700)] or [""]
            for fragment in fragments:
                if n + len(fragment) > 1550 and buffer:
                    text = "\n".join(part for _, part in buffer).strip()
                    if text:
                        chunks.append(Passage(path, page if kind == "page" else buffer[0][0],
                                              page if kind == "page" else buffer[-1][0],
                                              text, title, kind))
                    buffer = buffer[-1:]  # Minimal overlap; keep chunks small.
                    n = sum(len(value) for _, value in buffer)
                buffer.append((number, fragment))
                n += len(fragment) + 1
        if buffer:
            text = "\n".join(part for _, part in buffer).strip()
            if text:
                chunks.append(Passage(path, page if kind == "page" else buffer[0][0],
                                      page if kind == "page" else buffer[-1][0],
                                      text, title, kind))
        return chunks

    def _lexical_results(self, question: str) -> list[SearchHit]:
        """Existing document/heading/AGENTS-aware keyword ranking."""
        query = list(dict.fromkeys(words(question)))
        if not query or not self.passages:
            return []
        total_docs = max(1, len(self.doc_words))
        df = Counter()
        for term in query:
            df[term] = sum(term in token_set for token_set in self.doc_words.values())
        idf = {term: math.log(1 + (total_docs - df[term] + .5) / (df[term] + .5))
               for term in query}
        results: list[SearchHit] = []
        for passage in self.passages:
            route_count = Counter(words(self.routes.get(passage.path, "")))
            path_count = Counter(words(passage.path))
            head_count = Counter(words(passage.title))
            body_count = Counter(words(passage.content))
            score = 0.0
            match_count = 0
            for term in query:
                count = body_count[term]
                present = bool(count or head_count[term] or path_count[term] or route_count[term])
                if not present:
                    continue
                match_count += 1
                score += idf[term] * (
                    min(count, 3) * 1.0 +
                    min(head_count[term], 2) * 2.0 +
                    min(path_count[term], 2) * 4.0 +
                    min(route_count[term], 2) * 3.0
                )
            if match_count:
                score *= 1 + 0.17 * (match_count - 1)
                if passage.path.lower().endswith("agents.md") and "agent" not in query:
                    score *= 0.65
                results.append(SearchHit(passage, score))
        results.sort(key=lambda hit: hit.score, reverse=True)
        return results

    def _semantic_results(self, question: str) -> list[tuple[Passage, float]]:
        if (not self.semantic_enabled or self._semantic_vectors is None
                or self.semantic_status == "unavailable"
                or self._vectors_signature != self._semantic_signature):
            return []
        with self._encoder_lock:
            try:
                import numpy as np
                model = self._get_embedding_model()
                query_vectors = list(model.embed([question], batch_size=1))
                if len(query_vectors) != 1:
                    return []
                query_vec = self._normalize_vectors(query_vectors)[0]
                similarities = self._semantic_vectors @ query_vec
            except Exception as exc:
                # A failed local inference must never break trivia or Claude.
                self.semantic_last_error = f"Semantic query failed: {type(exc).__name__}"
                self.semantic_status = "unavailable"
                self._failed_signature = self._semantic_signature
                print(f"KB semantic query failed (keyword fallback): {exc}")
                return []
        # Avoid presenting entirely irrelevant documents simply because every
        # embedding search has a mathematical best match.
        ranked = np.argsort(-similarities)[:32]
        return [
            (self.passages[int(i)], float(similarities[i]))
            for i in ranked if similarities[i] >= 0.38
        ]

    def search(self, question: str, *, deep: bool = False) -> list[SearchHit]:
        """Hybrid rank fusion: lexical precision + semantic synonym matching.

        Use reciprocal rank fusion instead of comparing incompatible raw
        cosine-similarity and TF/IDF scores. This intentionally keeps filename
        and exact-word matches important for technical part numbers and rules.
        """
        lexical = self._lexical_results(question)
        semantic = self._semantic_results(question)
        candidates: dict[Passage, float] = defaultdict(float)
        for rank, hit in enumerate(lexical[:40], 1):
            candidates[hit.passage] += 1.45 / (24 + rank)
        for rank, (passage, similarity) in enumerate(semantic, 1):
            candidates[passage] += 1.65 * min(1.0, similarity / 0.65) / (24 + rank)
        if not candidates:
            return []
        ranked = sorted(candidates.items(), key=lambda item: item[1], reverse=True)
        if "agent" not in words(question) and any(
            p.path.rsplit("/", 1)[-1].lower() != "agents.md" for p, _ in ranked
        ):
            ranked = [(p, score) for p, score in ranked
                      if p.path.rsplit("/", 1)[-1].lower() != "agents.md"]
        max_passages = DEEP_MAX_PASSAGES if deep else MAX_PASSAGES
        max_files = DEEP_MAX_FILES if deep else MAX_FILES
        max_chars = DEEP_MAX_DOC_CHARS if deep else MAX_DOC_CHARS
        selected: list[SearchHit] = []
        file_hits: Counter[str] = Counter()
        used_chars = 0
        for passage, score in ranked:
            if file_hits[passage.path] >= (3 if deep else 2):
                continue
            if passage.path not in file_hits and len(file_hits) >= max_files:
                continue
            estimate = len(passage.content) + len(passage.path) + 100
            if used_chars + estimate > max_chars:
                continue
            selected.append(SearchHit(passage, score))
            file_hits[passage.path] += 1
            used_chars += estimate
            if len(selected) >= max_passages:
                break
        return selected

    async def search_async(self, question: str, *, deep: bool = False) -> list[SearchHit]:
        """Offload local ONNX inference so Discord reaction events stay responsive."""
        return await asyncio.to_thread(self.search, question, deep=deep)

    def excerpts(self, hits: list[SearchHit]) -> str:
        return "\n\n".join(
            f"[{i}] {hit.passage.path} ({hit.passage.location})\n{hit.passage.content}"
            for i, hit in enumerate(hits, 1)
        )

    def source_links(self, hits: list[SearchHit]) -> list[str]:
        seen = set()
        links = []
        for hit in hits:
            p = hit.passage
            if p.path in seen:
                continue
            seen.add(p.path)
            encoded = quote(p.path, safe="/")
            base = f"https://github.com/Mean-Green-Racing/fsae-kb/blob/{self.commit}/{encoded}"
            if p.source_kind == "lines":
                base += f"#L{p.start}-L{p.end}"
            elif p.source_kind == "page":
                base += f"#page={p.start}"
            links.append(f"[{p.path}]({base})")
        return links
